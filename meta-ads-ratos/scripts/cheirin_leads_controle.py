# -*- coding: utf-8 -*-
"""Planilha de controle de leads da Cheirin Bao MT (PINO SP, 14 Pinos e teste A/B de horario).

Uso:  python cheirin_leads_controle.py
Agendado no Task Scheduler ("CheirinLeadsControle-Mileny"), 2x ao dia.

Regrava a planilha inteira a cada execucao (idempotente), desde DESDE.
LGPD: NAO grava nome, e-mail, telefone nem respostas do formulario. So o id do lead
(texto, para cruzar com o CRM), data/hora e de onde o lead veio.
"""
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.expanduser("~"), "OneDrive", ".claude", "skills", "meta-ads-ratos")
ENV = os.path.join(BASE, ".env")
SAIDA = os.path.join(BASE, "relatorios", "controle", "cheirin-leads-controle.xlsx")
TOK = [l.split("=", 1)[1].strip().strip('"').strip("'") for l in open(ENV, encoding="utf-8")
       if l.strip().startswith("META_ADS_TOKEN")][0]
API = "https://graph.facebook.com/v21.0/"
BRT = timezone(timedelta(hours=-3))
DESDE = datetime(2026, 9, 1, tzinfo=BRT)

CONJUNTOS = {
    "120257337905370212": "PINO SP",
    "120257513484400212": "14 Pinos Lojas",
    "120257658661950212": "07A Teste - horarios Eduardo",
    "120257658674640212": "07B Teste - dia todo",
    "120256940806870212": "01 Leads Quente + Geo",
}
STATUS_ADS = ["ACTIVE", "PAUSED", "ADSET_PAUSED", "CAMPAIGN_PAUSED", "ARCHIVED", "WITH_ISSUES", "IN_PROCESS"]
DIAS = ["segunda", "terca", "quarta", "quinta", "sexta", "sabado", "domingo"]
# Horarios do Eduardo (isoweekday 1=seg ... 7=dom)
JANELAS = {1: [(6, 9), (18, 22)], 2: [(6, 9), (18, 22)], 3: [(6, 9), (18, 22)], 4: [(6, 9), (18, 22)],
           5: [(0, 24)], 6: [(8, 19)], 7: [(8, 16)]}


def faixa(h):
    if h < 6:
        return "0h-6h madrugada"
    if h < 9:
        return "6h-9h manha"
    if h < 12:
        return "9h-12h"
    if h < 15:
        return "12h-15h almoco"
    if h < 18:
        return "15h-18h"
    if h < 22:
        return "18h-22h noite"
    return "22h-24h"


def get(path, params):
    params = dict(params, access_token=TOK)
    url = API + path + "?" + urllib.parse.urlencode(params)
    out = []
    while url:
        for espera in (30, 60, 120, None):
            try:
                with urllib.request.urlopen(url, timeout=180) as r:
                    d = json.load(r)
                break
            except urllib.error.HTTPError as e:
                body = e.read().decode("utf-8", "replace")
                if espera and any(c in body for c in ('"code":17,', '"code":613', '"code":4,', '"code":80004', '"code":2,')):
                    time.sleep(espera)
                    continue
                raise RuntimeError("%s -> %s" % (path, body[:300]))
        out += d.get("data", [])
        url = d.get("paging", {}).get("next")
    return out


def coletar():
    linhas = []
    ts0 = int(DESDE.timestamp()) - 1
    hoje = datetime.now(BRT).strftime("%Y-%m-%d")
    for sid, rotulo in CONJUNTOS.items():
        ads = get(sid + "/ads", {"fields": "id,name,campaign{name},created_time", "effective_status": json.dumps(STATUS_ADS), "limit": 100})
        # So busca lead de anuncio que entregou no periodo (ou foi criado nele): anuncio sem impressao nao gera lead
        # pago, e varrer os arquivados antigos levava 18 min (16/09).
        entregou = {r["ad_id"] for r in get(sid + "/insights", {
            "fields": "ad_id,impressions", "level": "ad", "limit": 500,
            "time_range": json.dumps({"since": DESDE.strftime("%Y-%m-%d"), "until": hoje}),
            "filtering": json.dumps([{"field": "ad.effective_status", "operator": "IN", "value": STATUS_ADS}])})
            if int(r.get("impressions", 0)) > 0}
        ads = [a for a in ads if a["id"] in entregou or a.get("created_time", "") >= DESDE.strftime("%Y-%m-%d")]
        for ad in ads:
            leads = get(ad["id"] + "/leads", {"fields": "id,created_time,form_id,is_organic,platform",
                                               "filtering": json.dumps([{"field": "time_created", "operator": "GREATER_THAN", "value": ts0}]),
                                               "limit": 500})
            for ld in leads:
                dt = datetime.strptime(ld["created_time"], "%Y-%m-%dT%H:%M:%S%z").astimezone(BRT)
                dentro = any(a <= dt.hour < b for a, b in JANELAS[dt.isoweekday()])
                linhas.append({
                    "data": dt.strftime("%d/%m/%Y"), "hora": dt.strftime("%H:%M"), "dia": DIAS[dt.weekday()],
                    "faixa": faixa(dt.hour), "eduardo": "dentro" if dentro else "fora",
                    "conjunto": rotulo, "campanha": (ad.get("campaign") or {}).get("name", ""),
                    "anuncio": ad.get("name", ""), "formulario": str(ld.get("form_id", "")),
                    "lead_id": str(ld["id"]), "organico": "sim" if ld.get("is_organic") else "nao",
                    "plataforma": ld.get("platform", ""), "_dt": dt,
                })
            time.sleep(0.3)
    linhas.sort(key=lambda x: x["_dt"], reverse=True)
    return linhas


def cabecalho(ws, nomes, larguras):
    ws.append(nomes)
    for i, w in enumerate(larguras, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5C4B")
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def gravar(linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    cols = [("Data", "data", 12), ("Hora", "hora", 8), ("Dia", "dia", 10), ("Faixa", "faixa", 16),
            ("Horario Eduardo", "eduardo", 15), ("Conjunto", "conjunto", 28), ("Campanha", "campanha", 36),
            ("Anuncio", "anuncio", 38), ("Formulario", "formulario", 18), ("ID do lead (CRM)", "lead_id", 20),
            ("Organico", "organico", 10), ("Plataforma", "plataforma", 11)]
    cabecalho(ws, [c[0] for c in cols], [c[2] for c in cols])
    for l in linhas:
        ws.append([l[c[1]] for c in cols])
    for row in ws.iter_rows(min_row=2, min_col=10, max_col=10):
        for c in row:
            c.number_format = "@"
    ws.auto_filter.ref = ws.dimensions

    # Resumo por dia x conjunto
    rs = wb.create_sheet("Resumo por dia")
    conj = list(CONJUNTOS.values())
    cabecalho(rs, ["Data", "Dia"] + conj + ["Total", "Dentro do horario Eduardo", "Fora"],
              [12, 10] + [16] * len(conj) + [8, 14, 8])
    por_dia = defaultdict(lambda: defaultdict(int))
    dias_nome = {}
    for l in linhas:
        k = l["_dt"].strftime("%Y-%m-%d")
        dias_nome[k] = (l["data"], l["dia"])
        por_dia[k][l["conjunto"]] += 1
        por_dia[k]["_" + l["eduardo"]] += 1
    for k in sorted(por_dia, reverse=True):
        v = por_dia[k]
        tot = sum(v[c] for c in conj)
        rs.append(list(dias_nome[k]) + [v[c] for c in conj] + [tot, v["_dentro"], v["_fora"]])

    # Resumo por faixa de horario x conjunto
    rf = wb.create_sheet("Resumo por faixa")
    faixas = ["0h-6h madrugada", "6h-9h manha", "9h-12h", "12h-15h almoco", "15h-18h", "18h-22h noite", "22h-24h"]
    cabecalho(rf, ["Faixa"] + conj + ["Total"], [18] + [16] * len(conj) + [8])
    por_faixa = defaultdict(lambda: defaultdict(int))
    for l in linhas:
        por_faixa[l["faixa"]][l["conjunto"]] += 1
    for f in faixas:
        rf.append([f] + [por_faixa[f][c] for c in conj] + [sum(por_faixa[f][c] for c in conj)])

    info = wb.create_sheet("Sobre")
    for t in ["Planilha de controle - Cheirin Bao MT (act_884652616778712)",
              "Atualizada em %s (horario de Brasilia)" % datetime.now(BRT).strftime("%d/%m/%Y %H:%M"),
              "Leads desde %s, pela hora do cadastro no formulario da Meta." % DESDE.strftime("%d/%m/%Y"),
              "Conjuntos: " + ", ".join(CONJUNTOS.values()),
              "Horarios do Eduardo: seg-qui 6h-9h e 18h-22h; sexta dia todo; sabado 8h-19h; domingo 8h-16h.",
              "Sem nome, e-mail ou telefone (LGPD). Use o ID do lead para cruzar com o CRM.",
              "A planilha e regravada inteira a cada atualizacao: anotacoes feitas aqui se perdem."]:
        info.append([t])
    info.column_dimensions["A"].width = 100

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    tmp = SAIDA + ".tmp.xlsx"
    wb.save(tmp)
    try:
        os.replace(tmp, SAIDA)
        return SAIDA
    except PermissionError:
        alt = SAIDA.replace(".xlsx", "-%s.xlsx" % datetime.now(BRT).strftime("%Y%m%d-%H%M"))
        os.replace(tmp, alt)
        return alt


def main():
    linhas = coletar()
    caminho = gravar(linhas)
    cont = defaultdict(int)
    for l in linhas:
        cont[l["conjunto"]] += 1
    print("OK %d leads -> %s | %s" % (len(linhas), caminho, dict(cont)), flush=True)
    if caminho != SAIDA:
        print("AVISO: a planilha principal estava aberta no Excel; gravei uma copia com data.", file=sys.stderr, flush=True)


if __name__ == "__main__":
    main()
